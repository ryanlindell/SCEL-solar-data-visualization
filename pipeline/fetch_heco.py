"""Fetch Hawaiian Electric's published curtailment data (two Excel workbooks) and read the numbers
out of them. This module only fetches and parses - it never touches the database.

"Curtailment" means the utility told a wind farm or utility-scale solar plant to produce less than it
could have, usually because the grid could not use all of it. Hawaiian Electric reports the energy
that was curtailed and the energy it took ("delivered"), so potential = delivered + curtailed.

The workbooks are made for people, not programs, so the parsers look rows and headings up by their
labels and fail loudly if the layout changes, instead of quietly reading the wrong cells.
"""

import io
import re
import time

import requests
from openpyxl import load_workbook

from .config import CurtailmentSource, HecoFiles
from .http import MAX_ATTEMPTS

# Row labels in the totals workbook (matched by prefix) -> the name we store them under.
TOTALS_ROWS = {
    "2. mwh taken from curtailable": "delivered_mwh",
    "1. mwh curtailed": "curtailed_mwh",
    "3. mwh taken from firm": "firm_mwh",
    "4. mwh taken from uncurtailable": "distributed_mwh",
}
# Row labels in the by-reason workbook -> the name we store them under.
REASON_ROWS = {
    "oversupply": "oversupply_mwh",
    "system constraint": "system_constraint_mwh",
    "facility requested": "facility_requested_mwh",
    "total": "total_mwh",
}


class HecoError(Exception):
    """Hawaiian Electric's workbook could not be downloaded or did not look as expected."""


def normalize(text) -> str:
    """Lower-case, and ignore every kind of apostrophe/okina ('O‘ahu' == 'Oʻahu' == "O'ahu")."""
    return re.sub(r"[‘’ʻ'`]", "", str(text or "")).strip().lower()


def to_number(value) -> float | None:
    """A cell as a number, or None for blanks and markers such as 'NA'. (Some cells hold '0' as text.)"""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def parse_period(header) -> tuple[str, str] | None:
    """Column heading -> (frequency, period). 'Q1 2024' -> ('quarterly', '2024-Q1');
    2024 or '2024*' -> ('annual', '2024'). Anything else -> None."""
    text = str(header).strip() if header is not None else ""
    m = re.fullmatch(r"Q([1-4])\s+(\d{4})", text)
    if m:
        return "quarterly", f"{m.group(2)}-Q{m.group(1)}"
    m = re.fullmatch(r"(\d{4})\*{0,3}", text)
    if m:
        return "annual", m.group(1)
    return None


def parse_quarterly_totals(rows: list[tuple]) -> list[dict]:
    """The quarterly block of one island's sheet in the totals workbook.

    Returns rows: {"period": "2024-Q1", "series": "delivered_mwh", "value": 123.4}.
    """
    def is_quarterly_header(row: tuple) -> bool:
        parsed = parse_period(row[1]) if len(row) > 1 else None
        return row[0] == "B" and parsed is not None and parsed[0] == "quarterly"

    header_idx = next((i for i, r in enumerate(rows) if r and is_quarterly_header(r)), None)
    if header_idx is None:
        raise HecoError("Unexpected layout: no quarterly block (a row starting with 'B' and 'Q1 ...') found.")
    header = rows[header_idx]

    found: dict[str, tuple] = {}
    for row in rows[header_idx + 1 :]:
        label = normalize(row[0]) if row and row[0] is not None else ""
        if not label:
            break  # the block ends at the first blank label
        for prefix, name in TOTALS_ROWS.items():
            if label.startswith(prefix):
                found[name] = row
    missing = set(TOTALS_ROWS.values()) - set(found)
    if missing:
        raise HecoError(f"Unexpected layout: quarterly rows not found: {sorted(missing)}")

    out = []
    for col, head in enumerate(header):
        parsed = parse_period(head)
        if not parsed or parsed[0] != "quarterly":
            continue
        for name, row in found.items():
            value = to_number(row[col]) if col < len(row) else None
            if value is not None:
                out.append({"period": parsed[1], "series": name, "value": value})
    return out


def parse_by_reason(rows: list[tuple], section: str) -> list[dict]:
    """Curtailed energy by reason for one island, from the by-reason workbook.

    The sheet holds several blocks stacked vertically, each headed by a letter in the second column
    ('A' = annual Oʻahu, 'B' = quarterly Oʻahu, 'C' = annual Maui, ...). Each block has a heading row
    naming the island, then one row per reason. Returns rows:
    {"frequency": "quarterly", "period": "2024-Q1", "series": "oversupply_mwh", "value": 12.3}.
    """
    wanted = normalize(section)
    out: list[dict] = []
    i = 0
    while i < len(rows):
        row = rows[i]
        is_header = (
            len(row) > 2 and isinstance(row[1], str) and re.fullmatch(r"[A-Z]", row[1].strip() or "") is not None
            and parse_period(row[2]) is not None
        )
        if not is_header:
            i += 1
            continue
        header = row
        # The island heading follows, possibly after a row of "Annual" markers.
        j = i + 1
        while j < len(rows) and len(rows[j]) > 1 and normalize(rows[j][1]) in ("", "annual") and not _has_label(rows[j]):
            j += 1
        if j < len(rows) and len(rows[j]) > 1 and normalize(rows[j][1]) == wanted:
            k = j + 1
            while k < len(rows) and len(rows[k]) > 1 and normalize(rows[k][1]) in REASON_ROWS:
                for col, head in enumerate(header):
                    parsed = parse_period(head) if col >= 2 else None
                    value = to_number(rows[k][col]) if parsed and col < len(rows[k]) else None
                    if value is not None:
                        out.append(
                            {
                                "frequency": parsed[0],
                                "period": parsed[1],
                                "series": REASON_ROWS[normalize(rows[k][1])],
                                "value": value,
                            }
                        )
                k += 1
        i += 1
    if not out:
        raise HecoError(f"Unexpected layout: no by-reason rows found for '{section}'.")
    return out


def _has_label(row: tuple) -> bool:
    return len(row) > 1 and row[1] is not None and normalize(row[1]) not in ("", "annual")


def _download(url: str, session: requests.Session) -> bytes:
    last_problem = "unknown error"
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            resp = session.get(url, timeout=60)
        except requests.RequestException as exc:
            last_problem = f"network error ({type(exc).__name__})"
        else:
            if resp.status_code == 200:
                return resp.content
            if resp.status_code == 429 or resp.status_code >= 500:
                last_problem = f"HTTP {resp.status_code}"
            else:
                raise HecoError(f"HTTP {resp.status_code} downloading {url.rsplit('/', 1)[-1]}")
        if attempt < MAX_ATTEMPTS:
            time.sleep(2**attempt)
    raise HecoError(f"Giving up on {url.rsplit('/', 1)[-1]} after {MAX_ATTEMPTS} attempts: {last_problem}")


def _sheet_rows(content: bytes, sheet: str | None) -> list[tuple]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a bad file
        raise HecoError(f"Could not read the Excel file ({type(exc).__name__}).") from None
    if sheet is None:
        ws = wb.worksheets[0]
    elif sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        raise HecoError(f"Sheet '{sheet}' not found; the workbook has {wb.sheetnames}.")
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def fetch_curtailment(
    source: CurtailmentSource, files: HecoFiles, session: requests.Session | None = None
) -> dict[str, list[dict]]:
    """Download both workbooks and return {"quarterly_totals": [...], "by_reason": [...]}."""
    session = session or requests.Session()
    totals = _sheet_rows(_download(files.base_url + files.totals_file, session), source.totals_sheet)
    reasons = _sheet_rows(_download(files.base_url + files.by_reason_file, session), None)
    return {
        "quarterly_totals": parse_quarterly_totals(totals),
        "by_reason": parse_by_reason(reasons, source.by_reason_section),
    }
